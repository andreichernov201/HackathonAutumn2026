"""Preview behavior against all seven real source templates."""
import hashlib
import io
import os
import tempfile
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from lxml import etree

os.environ.setdefault('REGULA_DATA_DIR', tempfile.mkdtemp(prefix='regula-preview-tests-'))
os.environ.setdefault('REGULA_SESSION_SECRET', 'preview-test-secret-at-least-32-characters')

from api import main
from api.preview_markup import FILLED_FILL, MISSING_FILL
from api.service import CATALOG, DocumentsService, SOURCE_DIR

IDS = [item[0] for item in CATALOG]
MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


@pytest.fixture
def api(tmp_path, monkeypatch):
    service = DocumentsService(tmp_path)
    monkeypatch.setattr(main, 'service', service)
    client = TestClient(main.app)
    session = client.get('/api/session').json()
    return client, {'X-CSRF-Token': session['csrfToken']}, service


def sample(specs):
    fields = {spec.key: spec.options[0] if spec.options else '2026-12-31'
              if spec.inputType == 'date' else f'Значение {spec.key}'
              for spec in specs if spec.required}
    if 'appointees' in fields:
        fields['appointees'] = '\n'.join(f'Отдел {i} | Специалист {i}' for i in range(4))
    return fields


def runs(document):
    return list(document.element.iter(qn('w:r')))


def marked(document, fill):
    return [run for run in runs(document) if any(shd.get(qn('w:fill')) == fill
            for shd in run.iter(qn('w:shd')))]


def text(document):
    return '\n'.join(''.join(node.text or '' for node in paragraph.iter(qn('w:t')))
                     for paragraph in document.element.iter(qn('w:p')))


def state(service):
    with service.db() as db:
        rows = {table: [tuple(row) for row in db.execute(f'SELECT * FROM {table}')]
                for table in ('projects', 'artifacts', 'templates')}
    return rows, {str(path.relative_to(service.root)): hashlib.sha256(path.read_bytes()).hexdigest()
                  for path in service.root.rglob('*') if path.is_file()}


@pytest.mark.parametrize('template_id', IDS)
@pytest.mark.parametrize('kind', ['empty', 'partial', 'full'])
def test_seven_templates_preview_highlights_and_final_is_clean(api, template_id, kind):
    client, headers, service = api
    template = service.template(template_id)
    source_path = SOURCE_DIR / template.sourceFilename
    source_bytes = source_path.read_bytes()
    original = Document(io.BytesIO(source_bytes))
    project = client.post('/api/document-projects', headers=headers, json={'templateId': template_id}).json()
    fields = {} if kind == 'empty' else {'organization': 'РАСПОРЯЖЕНИЕ'} if kind == 'partial' else sample(template.fields)
    if kind == 'full':
        multiline = next((spec.key for spec in template.fields if spec.required and spec.multiline and spec.key != 'appointees'), None)
        if multiline:
            fields[multiline] = 'Первая строка\nВторая строка'
    before = state(service)
    response = client.post(f"/api/document-projects/{project['projectId']}/preview-file", headers=headers, json={'fields': fields})
    assert response.status_code == 200, response.text
    assert response.headers['content-type'] == MIME
    assert response.headers['cache-control'] == 'private, no-store'
    assert response.headers['x-content-type-options'] == 'nosniff'
    assert response.headers['content-disposition'].startswith('inline')
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert archive.testzip() is None
        assert '[Content_Types].xml' in archive.namelist()
        with zipfile.ZipFile(io.BytesIO(source_bytes)) as source_zip:
            for name in source_zip.namelist():
                if name.startswith(('word/header', 'word/footer')) or name in ('word/styles.xml', 'word/numbering.xml'):
                    assert etree.tostring(etree.fromstring(archive.read(name)), method='c14n') == etree.tostring(etree.fromstring(source_zip.read(name)), method='c14n')
    document = Document(io.BytesIO(response.content))
    content = text(document)
    assert '\ue000' not in content and '\ue001' not in content and '\ue002' not in content
    assert bool(marked(document, MISSING_FILL)) == (kind != 'full')
    assert bool(marked(document, FILLED_FILL)) == (kind != 'empty')
    if kind == 'partial':
        assert document.paragraphs[0].text == 'РАСПОРЯЖЕНИЕ'
        # An identical fixed word elsewhere is not accidentally highlighted.
        assert not any(shd.get(qn('w:fill')) == FILLED_FILL
                       for shd in document.paragraphs[1]._p.iter(qn('w:shd')))
    assert 'ДАТА / РЕГИСТРАЦИОННЫЙ НОМЕР' in content and 'ВИЗУАЛИЗАЦИЯ ЭП' in content
    assert 'МОДУЛЬНЫЙ ДОКУМЕНТ' not in content and 'УТВЕРЖДАЮ' not in content
    assert 'Примечания по заполнению' not in content
    assert document.sections[0]._sectPr.xml == original.sections[0]._sectPr.xml
    assert len(document.tables) == len(original.tables) - 1
    for actual, expected in zip(document.tables, original.tables[1:]):
        assert actual._tbl.tblPr.xml == expected._tbl.tblPr.xml
        assert actual._tbl.tblGrid.xml == expected._tbl.tblGrid.xml
    assert source_path.read_bytes() == source_bytes
    assert state(service) == before
    if kind != 'full':
        assert client.post(f"/api/document-projects/{project['projectId']}/generate", headers=headers).status_code == 422
    else:
        updated = client.patch(f"/api/document-projects/{project['projectId']}", headers=headers,
                               json={'fields': fields, 'expectedVersion': 1})
        assert updated.status_code == 200 and updated.json()['validation']['ready']
        generated = client.post(f"/api/document-projects/{project['projectId']}/generate", headers=headers)
        assert generated.status_code == 200, generated.text
        artifact = generated.json()['artifacts'][0]
        final = client.get(f"/api/document-projects/{project['projectId']}/artifacts/{artifact['artifactId']}/download")
        final_doc = Document(io.BytesIO(final.content))
        assert not marked(final_doc, MISSING_FILL) and not marked(final_doc, FILLED_FILL)
        assert 'Укажите:' not in text(final_doc) and '\ue000' not in text(final_doc)


@pytest.mark.parametrize('fields', [
    {'unknown': 'value'}, {'organization': 'x' * 4001}, {'organization': 5},
    {'organization': None}, {'organization': ['value']}, {'organization': 'test\x00value'},
    {'deadline': '2026-02-31'}, {'subject': '[Подсказка]'},
])
def test_invalid_preview_values_rejected_without_mutation(api, fields):
    client, headers, service = api
    project = client.post('/api/document-projects', headers=headers, json={'templateId': 'working-group-order'}).json()
    before = state(service)
    response = client.post(f"/api/document-projects/{project['projectId']}/preview-file", json={'fields': fields}, headers=headers)
    assert response.status_code == 422
    assert state(service) == before


def test_access_csrf_ids_and_parallel_isolation(api):
    client, headers, service = api
    project = client.post('/api/document-projects', headers=headers, json={'templateId': 'working-group-order'}).json()
    url = f"/api/document-projects/{project['projectId']}/preview-file"
    assert client.post(url, json={'fields': {}}).status_code == 403
    assert client.post('/api/document-projects/invalid/preview-file', json={'fields': {}}, headers=headers).status_code == 404
    other = TestClient(main.app)
    other_headers = {'X-CSRF-Token': other.get('/api/session').json()['csrfToken']}
    assert other.post(url, headers=other_headers, json={'fields': {}}).status_code == 404
    before = state(service)

    def request(index):
        response = client.post(url, headers=headers, json={'fields': {'organization': f'Организация-{index:02d}'}})
        assert response.status_code == 200
        content = text(Document(io.BytesIO(response.content)))
        assert f'Организация-{index:02d}' in content
        assert sum(f'Организация-{j:02d}' in content for j in range(12)) == 1

    with ThreadPoolExecutor(max_workers=4) as pool:
        list(pool.map(request, range(12)))
    request(1)
    assert state(service) == before


def test_optional_conditions_dates_and_tables(api):
    client, headers, service = api
    for template_id, fields, expected in (
        ('project-committee-order', {'projectForm': 'проектам', 'projects': 'Проект А', 'cancelNumber': '42'}, '3. Контроль'),
        ('project-committee-order', {}, '2. Контроль'),
        ('project-committee-order', {'cancelNumber': '  '}, '2. Контроль'),
        ('handover-order', {'source': 'от начальника Управления', 'deadline': '2026-12-31'}, 'от начальника Управления'),
        ('local-act-order', {'cancelOrder': 'приказ № 12'}, '5. Контроль'),
        ('records-officers-order', {'appointees': 'Отдел А | Иванов\nОтдел Б | Петров'}, 'Петров'),
    ):
        project = client.post('/api/document-projects', headers=headers, json={'templateId': template_id}).json()
        response = client.post(f"/api/document-projects/{project['projectId']}/preview-file", headers=headers, json={'fields': fields})
        assert response.status_code == 200, response.text
        content = text(Document(io.BytesIO(response.content)))
        assert expected in content
        assert 'от от начальника' not in content
        if template_id == 'handover-order':
            assert '31» декабря 2026 г.' in content
    for template_id, fields in (
        ('project-committee-order', {'projectForm': 'invalid'}),
        ('records-officers-order', {'appointees': 'нет разделителя'}),
        ('internal-audit-order', {'startDate': '2026-12-31', 'endDate': '2026-01-01'}),
    ):
        project = client.post('/api/document-projects', headers=headers, json={'templateId': template_id}).json()
        assert client.post(f"/api/document-projects/{project['projectId']}/preview-file", headers=headers, json={'fields': fields}).status_code == 422


@pytest.mark.parametrize('failure', ['missing', 'corrupt', 'unavailable', 'version', 'structure'])
def test_source_errors_do_not_expose_details_or_change_project(api, failure):
    client, headers, service = api
    project = client.post('/api/document-projects', headers=headers, json={'templateId': 'working-group-order'}).json()
    source = service.root / 'templates' / f"{service.template('working-group-order').sourceSha256}.docx"
    if failure == 'missing':
        source.unlink()
    elif failure == 'corrupt':
        source.write_bytes(b'invalid DOCX at /private/internal/source')
    elif failure == 'structure':
        doc = Document(source)
        doc.paragraphs[5].text = 'Changed source'
        doc.save(source)
    else:
        with service.db() as db:
            if failure == 'unavailable':
                db.execute("UPDATE templates SET status='awaiting_review' WHERE id='working-group-order'")
            else:
                db.execute("UPDATE templates SET version='changed' WHERE id='working-group-order'")
    before = state(service)
    response = client.post(f"/api/document-projects/{project['projectId']}/preview-file", headers=headers, json={'fields': {}})
    assert response.status_code == 422
    assert '/private/' not in response.text and str(service.root) not in response.text
    assert state(service) == before
