"""Cross-service contracts on one application, database and signed session."""
import hashlib
import io
import json
import time
from collections import Counter

import pytest
from docx import Document
from fastapi import FastAPI, File, UploadFile
from fastapi.testclient import TestClient

from api import main
from api.config import session_secret
from api.knowledge.seed import import_case_sources
from api.knowledge.service import KnowledgeService
from api.limits import MAX_REQUEST_BYTES, RequestSizeLimitMiddleware
from api.service import DocumentsService, SOURCE_DIR
from api.session import SESSION_MAX_AGE
from api.tool_registry import ToolRegistry

TOKEN = 'integration-maintenance-key-at-least-32-characters'


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    documents = DocumentsService(tmp_path)
    knowledge = KnowledgeService(tmp_path, documents.db_path)
    import_case_sources(knowledge)
    monkeypatch.setattr(main, 'service', documents)
    # The existing knowledge router closes over the original service instance.
    monkeypatch.setattr(main.knowledge_service, 'root', knowledge.root)
    monkeypatch.setattr(main.knowledge_service, 'db_path', knowledge.db_path)
    monkeypatch.setenv('REGULA_KNOWLEDGE_ADMIN_TOKEN', TOKEN)
    client = TestClient(main.app)
    session = client.get('/api/session').json()
    client.headers['X-CSRF-Token'] = session['csrfToken']
    return client, session, documents, knowledge


def invoke(client, name, **arguments):
    return client.post(f'/api/tools/{name}/invoke', json={'arguments': arguments})


def test_one_session_database_and_tools_from_search_to_generated_document(workspace):
    client, session, documents, knowledge = workspace
    assert documents.db_path == knowledge.db_path
    routes = Counter((method, route.path) for route in main.app.routes for method in getattr(route, 'methods', []))
    assert all(count == 1 for count in routes.values())
    registry = client.get('/api/tools').json()
    assert len(registry) == len({tool['name'] for tool in registry}) == 16
    assert all(tool['inputSchema'].get('additionalProperties') is False for tool in registry)
    assert all(tool['path'] == f"/api/tools/{tool['name']}/invoke" for tool in registry)
    source_bytes = {path.name: path.read_bytes() for path in SOURCE_DIR.glob('*.docx')}
    created = invoke(client, 'create_document_project', templateId='working-group-order').json()
    project_id = created['projectId']
    login = client.post('/api/knowledge/admin-session', json={'token': TOKEN})
    assert login.status_code == 200
    assert f'Max-Age={SESSION_MAX_AGE}' in login.headers['set-cookie']
    assert client.get('/api/session').json() == session  # Admin grant preserves device + CSRF.
    assert client.get('/api/knowledge/session').json()['role'] == 'knowledge_admin'
    catalog = invoke(client, 'list_knowledge_documents').json()
    assert catalog['total'] == 3
    found = invoke(client, 'search_knowledge', q='Когда присваивается регистрационный номер проекту распоряжения?').json()['items'][0]
    assert found['section']['sectionNumber'] == '2.1'
    citation = found['citation']
    point = invoke(client, 'get_knowledge_section', sectionId=citation['sectionId']).json()
    assert point['source']['checksum'] == citation['checksum']
    original = client.get(citation['originalUrl'])
    assert original.status_code == 200
    assert hashlib.sha256(original.content).hexdigest() == citation['checksum']
    assert invoke(client, 'verify_citation', **{key: citation[key] for key in ('documentId', 'versionId', 'sectionId', 'excerpt')}).json()['verified']
    fields = {spec.key: '2026-12-31' if spec.inputType == 'date' else f'Интеграция {spec.key}' for spec in documents.template('working-group-order').fields if spec.required}
    updated = invoke(client, 'update_document_fields', projectId=project_id, expectedVersion=1, fields=fields).json()
    assert updated['validation']['ready']
    preview = client.post(f'/api/document-projects/{project_id}/preview', json={'fields': fields})
    assert preview.status_code == 200 and preview.json()['preview']
    assert invoke(client, 'validate_document_project', projectId=project_id).json()['validation']['ready']
    generated = invoke(client, 'generate_document_docx', projectId=project_id).json()
    artifact = generated['artifacts'][0]
    link = invoke(client, 'get_document_artifact', projectId=project_id, artifactId=artifact['artifactId']).json()['downloadUrl']
    response = client.get(link)
    assert hashlib.sha256(response.content).hexdigest() == artifact['sha256']
    assert Document(io.BytesIO(response.content)).tables
    for path in SOURCE_DIR.glob('*.docx'):
        assert path.read_bytes() == source_bytes[path.name]
    assert client.post(f"/api/knowledge/versions/{citation['versionId']}/index").status_code == 200
    assert client.get('/api/knowledge/sections/' + point['id']).json()['text'] == point['text']
    # Expiring the maintenance grant must downgrade the role without losing
    # the persistent device cookie or document ownership.
    signed = main.signer.loads(client.cookies.get('regula_session'))
    signed['knowledgeGrantedAt'] = time.time() - 9 * 60 * 60
    client.cookies.set('regula_session', main.signer.dumps(signed), domain='testserver.local', path='/')
    assert client.get('/api/knowledge/session').json()['role'] == 'reader'
    assert client.get('/api/session').json() == session
    logout = client.delete('/api/knowledge/admin-session')
    assert logout.status_code == 200
    assert f'Max-Age={SESSION_MAX_AGE}' in logout.headers['set-cookie']
    assert client.get('/api/session').json() == session
    assert invoke(client, 'get_document_project', projectId=project_id).json()['projectId'] == project_id
    assert client.get('/api/document-projects').json()[0]['projectId'] == project_id
    outsider = TestClient(main.app)
    outsider.headers['X-CSRF-Token'] = outsider.get('/api/session').json()['csrfToken']
    assert invoke(outsider, 'get_document_project', projectId=project_id).status_code == 404
    assert outsider.get(link).status_code == 404


def test_common_tool_registry_enforces_csrf_schema_and_acl(workspace):
    client, _, _, _ = workspace
    anonymous = TestClient(main.app)
    assert invoke(anonymous, 'search_knowledge', q='номер').status_code == 403
    assert invoke(client, 'run_sql', sql='SELECT 1').status_code == 404
    assert invoke(client, 'list_templates', role='knowledge_admin').status_code == 422
    assert invoke(client, 'create_document_project', templateId='working-group-order', owner='another').status_code == 422
    assert invoke(client, 'search_knowledge', q='номер', role='knowledge_admin').status_code == 422
    assert client.post('/api/knowledge/admin-session', json={'token': TOKEN}).status_code == 200
    restricted = client.post('/api/knowledge/documents', data={'metadata': json.dumps({'title': 'Закрытая инструкция', 'category': 'instruction', 'accessLevel': 'restricted'})}, files={'file': ('restricted.txt', '2.1 Закрытый тестовый пункт.'.encode(), 'text/plain')}).json()
    point = client.get(f"/api/knowledge/versions/{restricted['latestVersion']['id']}/sections").json()['items'][0]
    assert invoke(client, 'get_knowledge_section', sectionId=point['id']).status_code == 200
    assert client.delete('/api/knowledge/admin-session').status_code == 200
    assert invoke(client, 'get_knowledge_section', sectionId=point['id']).status_code == 404
    assert invoke(client, 'get_knowledge_document', documentId=restricted['id']).status_code == 404
    registry = ToolRegistry()
    registry.register([{'name': 'test'}], lambda *args: None)
    with pytest.raises(ValueError, match='duplicate_tool'):
        registry.register([{'name': 'test'}], lambda *args: None)


def test_single_appointee_visible_in_preview_and_docx(workspace):
    client, _, documents, _ = workspace
    template_id = 'records-officers-order'
    template = documents.template(template_id)
    fields = {spec.key: 'Интеграция ' + spec.key for spec in template.fields if spec.required}
    fields['appointees'] = 'Отдел интеграции | специалист И. И. Иванов'
    project = client.post('/api/document-projects', json={'templateId': template_id}).json()
    preview = client.post(f"/api/document-projects/{project['projectId']}/preview", json={'fields': fields}).json()
    text = '\n'.join(item['label'] + ' ' + item['value'] for item in preview['preview'])
    assert 'Отдел интеграции' in text and 'И. И. Иванов' in text
    pid = project['projectId']
    client.patch(f'/api/document-projects/{pid}', json={'expectedVersion': 1, 'fields': fields})
    generated = client.post(f'/api/document-projects/{pid}/generate')
    assert generated.status_code == 200, generated.text
    artifact = generated.json()['artifacts'][0]
    content = client.get(f"/api/document-projects/{pid}/artifacts/{artifact['artifactId']}/download").content
    table = Document(io.BytesIO(content)).tables[1]
    assert len(table.rows) == 2
    assert 'Отдел интеграции' in table.rows[1].cells[0].text
    assert 'И. И. Иванов' in table.rows[1].cells[1].text


def test_api_limits_declared_and_streamed_multipart(workspace):
    client, _, _, _ = workspace
    response = client.post('/api/knowledge/documents', headers={'Content-Length': str(MAX_REQUEST_BYTES + 1)})
    assert response.status_code == 413
    limited = FastAPI()
    limited.add_middleware(RequestSizeLimitMiddleware, max_bytes=512)

    @limited.post('/api/upload')
    async def upload(file: UploadFile = File()):
        return {'size': len(await file.read())}

    def multipart(size):
        yield b'--test\r\nContent-Disposition: form-data; name="file"; filename="test.txt"\r\nContent-Type: text/plain\r\n\r\n'
        yield b'a' * size
        yield b'\r\n--test--\r\n'

    with TestClient(limited) as stream_client:
        headers = {'Content-Type': 'multipart/form-data; boundary=test'}
        assert stream_client.post('/api/upload', headers=headers, content=multipart(10)).json() == {'size': 10}
        oversized = stream_client.post('/api/upload', headers=headers, content=multipart(600))
        assert oversized.status_code == 413, oversized.text


def test_session_configuration_fails_closed_in_production(monkeypatch):
    monkeypatch.delenv('REGULA_SESSION_SECRET', raising=False)
    monkeypatch.setenv('REGULA_ENV', 'development')
    assert len(session_secret()) >= 32 and session_secret() != session_secret()
    monkeypatch.setenv('REGULA_ENV', 'production')
    with pytest.raises(RuntimeError, match='REGULA_SESSION_SECRET'):
        session_secret()
    for value in ('local-development-only-change-me', 'replace-with-long-random-secret', 'short'):
        monkeypatch.setenv('REGULA_SESSION_SECRET', value)
        with pytest.raises(RuntimeError):
            session_secret()
    monkeypatch.setenv('REGULA_SESSION_SECRET', 'a-private-operator-secret-at-least-32-characters')
    monkeypatch.setenv('REGULA_SECURE_COOKIE', '0')
    with pytest.raises(RuntimeError, match='REGULA_SECURE_COOKIE'):
        session_secret()
    monkeypatch.setenv('REGULA_SECURE_COOKIE', '1')
    assert session_secret().startswith('a-private-operator')
