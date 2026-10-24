"""Generated fixtures test mechanics, never replace the missing case originals."""
import hashlib
import io
import json
import os
from pathlib import Path
from uuid import uuid4

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient

from api.main import app
from api.knowledge.extract import MAX_BYTES, extract
from api.knowledge.models import ListArgs, Principal
from api.knowledge.service import KnowledgeService

TOKEN = 'test-knowledge-maintenance-key-at-least-32-characters'
DOCX_MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


@pytest.fixture
def corpus(tmp_path, monkeypatch):
    # The production router closes over its service, so swap only its storage paths.
    from api.main import knowledge_service
    isolated = KnowledgeService(tmp_path, tmp_path / 'test.sqlite3')
    monkeypatch.setattr(knowledge_service, 'db_path', isolated.db_path)
    monkeypatch.setattr(knowledge_service, 'root', isolated.root)
    monkeypatch.setenv('REGULA_KNOWLEDGE_ADMIN_TOKEN', TOKEN)
    return isolated


def device(admin=False):
    client = TestClient(app)
    session = client.get('/api/session').json()
    client.headers['X-CSRF-Token'] = session['csrfToken']
    if admin:
        assert client.post('/api/knowledge/admin-session', json={'token': TOKEN}).status_code == 200
    return client


def metadata(**values):
    return {'title': 'Учебная тестовая инструкция по документообороту', 'category': 'instruction', **values}


def upload(client, content=None, name='fixture.txt', meta=None, document_id=None, mime='text/plain'):
    # Invented test text; intentionally not a statement about case requirements.
    content = content if content is not None else 'Глава 1. Тестовые положения\n1.1 Тестовое требование об обработке документов.\n2.1 Регистрационный номер в этой тестовой фикстуре присваивается после тестовой проверки.'.encode()
    url = '/api/knowledge/documents' + (f'/{document_id}/versions' if document_id else '')
    return client.post(url, data={'metadata': json.dumps(meta or metadata())}, files={'file': (name, content, mime)})


def sections(client, doc):
    return client.get(f"/api/knowledge/versions/{doc['latestVersion']['id']}/sections").json()['items']


def invoke(client, name, arguments):
    return client.post(f'/api/knowledge/tools/{name}/invoke', json={'arguments': arguments})


def test_upload_catalog_full_text_phrase_and_unknown_requirement(corpus):
    admin = device(True)
    created = upload(admin)
    assert created.status_code == 201, created.text
    doc = created.json()
    reader = device()
    catalog = reader.get('/api/knowledge/documents').json()
    assert catalog['total'] == 1 and catalog['items'][0]['id'] == doc['id']
    assert doc['latestVersion']['status'] == 'unverified'
    assert doc['latestVersion']['versionLabel'] is None
    for query in ['инструкция по документообороту', '2.1', 'Когда присваивается регистрационный номер?', '"Регистрационный номер"']:
        response = reader.get('/api/knowledge/search', params={'q': query})
        assert response.status_code == 200, response.text
        assert response.json()['total'] > 0
        hit = response.json()['items'][0]
        assert hit['citation']['versionId'] == doc['latestVersion']['id']
        assert hit['citation']['sourceReference'].endswith(hit['section']['id'])
    unknown = reader.get('/api/knowledge/search', params={'q': 'квазикристаллическая телепортация'}).json()
    assert unknown['items'] == [] and 'Нет подтверждённого основания' in unknown['message']
    assert reader.get('/api/knowledge/documents', params={'status': 'active'}).json()['total'] == 0
    assert reader.get('/api/knowledge/search', params={'q': '" OR NEAR() --'}).status_code == 200
    assert reader.get('/api/knowledge/search', params={'q': '*' }).json()['total'] == 0
    assert reader.get('/api/knowledge/search', params={'q': 'тест', 'limit': 101}).status_code == 422
    assert reader.get('/api/knowledge/documents', params={'sort': 'title; DROP TABLE projects'}).status_code == 422
    page = reader.get('/api/knowledge/search', params={'q': 'тест', 'limit': 1, 'offset': 1}).json()
    assert page['limit'] == 1 and len(page['items']) == 1


def test_docx_matrix_rows_keep_required_conditional_columns(corpus):
    document = Document()
    document.add_heading('Внутренние документы', 1)
    table = document.add_table(rows=1, cols=4)
    for cell, value in zip(table.rows[0].cells, ['Номер', 'Документ', 'Обязательные (тест)', 'Условные (тест)']):
        cell.text = value
    for cell, value in zip(table.add_row().cells, ['6', 'Тест: рабочая группа для ИТ-проекта', 'Тестовое подразделение А', 'Тестовое подразделение Б: только для ИТ']):
        cell.text = value
    data = io.BytesIO(); document.save(data)
    admin = device(True)
    response = upload(admin, data.getvalue(), 'matrix.docx', metadata(title='Тестовая матрица', category='matrix'), mime=DOCX_MIME)
    assert response.status_code == 201, response.text
    doc = response.json()
    rows = sections(admin, doc)
    row = next(row for row in rows if row['sectionNumber'] == '6')
    point = admin.get('/api/knowledge/sections/' + row['id']).json()
    assert point['table'][0][2] == 'Обязательные (тест)'
    assert 'Условные (тест): Тестовое подразделение Б: только для ИТ' in point['text']
    assert 'Внутренние документы' in point['source']['locator'] and 'строка 6' in point['source']['locator']
    assert doc['latestVersion']['pageCount'] is None
    found = admin.get('/api/knowledge/search', params={'q': 'рабочая группа ИТ', 'category': 'matrix'}).json()
    assert any(item['section']['id'] == row['id'] for item in found['items'])
    original = admin.get(doc['latestVersion']['originalUrl'])
    assert original.content == data.getvalue()


def test_pdf_page_anchors_and_scanned_pdf_rejected(corpus):
    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), '2.1 Registration test requirement.')
    pdf.new_page().insert_text((72, 72), '2.2 Another test requirement.')
    content = pdf.tobytes(); pdf.close()
    admin = device(True)
    response = upload(admin, content, 'manual.pdf', mime='application/pdf')
    assert response.status_code == 201, response.text
    doc = response.json()
    assert doc['latestVersion']['pageCount'] == 2
    point = admin.get('/api/knowledge/sections/' + sections(admin, doc)[0]['id']).json()
    assert point['pageStart'] == point['pageEnd'] == 1
    assert point['source']['originalUrl'].endswith('#page=1')
    assert admin.get(doc['latestVersion']['originalUrl']).content == content
    blank = fitz.open(); blank.new_page()
    assert upload(admin, blank.tobytes(), 'scan.pdf', mime='application/pdf').json()['detail'] == 'no_text_layer'
    blank.close()


def test_editions_immutable_links_reindex_and_checksum(corpus):
    admin = device(True)
    doc = upload(admin).json()
    old_version = doc['latestVersion']
    old_sections = sections(admin, doc)
    old_point = admin.get('/api/knowledge/sections/' + old_sections[-1]['id']).json()
    new = upload(admin, '2.1 Вторая тестовая редакция: совершенно иное требование.'.encode(), meta=metadata(versionLabel='Тест 2'), document_id=doc['id'])
    assert new.status_code == 201, new.text
    assert new.json()['latestVersion']['status'] == 'unverified'
    editions = admin.get(f"/api/knowledge/documents/{doc['id']}/versions").json()
    assert editions['total'] == 2
    assert admin.get('/api/knowledge/sections/' + old_point['id']).json()['text'] == old_point['text']
    assert upload(admin, document_id=doc['id']).status_code == 409
    assert admin.get(old_version['originalUrl']).status_code == 200
    result = admin.post(f"/api/knowledge/versions/{old_version['id']}/index")
    assert result.status_code == 200
    assert admin.get(f"/api/knowledge/versions/{old_version['id']}/sections").json()['items'] == old_sections
    assert admin.get('/api/knowledge/search', params={'q': 'регистрационный', 'versionId': old_version['id']}).json()['total'] > 0
    assert admin.get('/api/knowledge/search', params={'q': 'регистрационный', 'versionId': new.json()['latestVersion']['id']}).json()['total'] == 0
    with corpus.db() as db:
        assert db.execute('SELECT COUNT(*) FROM knowledge_audit').fetchone()[0] == 3
        assert db.execute("SELECT COUNT(*) FROM knowledge_chunks WHERE embedding_status='not_configured'").fetchone()[0] > 0
    path = corpus.root / f"{old_version['id']}.txt"
    path.write_text('tampered')
    assert admin.get(old_version['originalUrl']).status_code == 409


def test_shared_ai_tools_chain_and_exact_citation_verification(corpus):
    admin = device(True)
    upload(admin)
    reader = device()
    tools = reader.get('/api/knowledge/tools').json()
    assert reader.get('/api/docs').status_code == 200
    assert '/api/knowledge/search' in reader.get('/api/openapi.json').json()['paths']
    assert len(tools) == 6
    assert all(tool['inputSchema']['additionalProperties'] is False for tool in tools)
    assert all(tool['sourceContentTrust'] == 'untrusted_data' for tool in tools)
    found = invoke(reader, 'search_knowledge', {'q': 'регистрационный номер'}).json()['items'][0]
    point = invoke(reader, 'get_knowledge_section', {'sectionId': found['section']['id']}).json()
    args = {key: point['source'][key] for key in ('documentId', 'versionId', 'sectionId')}
    verified = invoke(reader, 'verify_citation', {**args, 'excerpt': found['citation']['excerpt']})
    assert verified.status_code == 200 and verified.json()['verified']
    assert not invoke(reader, 'verify_citation', {**args, 'excerpt': 'Несуществующее требование'}).json()['verified']
    assert invoke(reader, 'verify_citation', {**args, 'versionId': str(uuid4()), 'excerpt': point['text']}).json()['reason'] == 'source_mismatch'
    assert invoke(reader, 'search_knowledge', {'q': 'тест', 'role': 'knowledge_admin'}).status_code == 422
    assert invoke(reader, 'run_sql', {'sql': 'SELECT * FROM documents'}).status_code == 404
    assert invoke(reader, 'get_knowledge_section', {'sectionId': '../../documents.sqlite3'}).status_code == 422
    assert len(invoke(reader, 'list_knowledge_documents', {}).json()['items']) == 1
    assert invoke(reader, 'get_knowledge_document', {'documentId': args['documentId']}).status_code == 200
    assert invoke(reader, 'list_knowledge_versions', {'documentId': args['documentId']}).json()['total'] == 1


def test_restricted_access_all_routes_and_tools_and_csrf(corpus):
    admin = device(True)
    doc = upload(admin, meta=metadata(accessLevel='restricted')).json()
    version = doc['latestVersion']; point = sections(admin, doc)[0]
    reader = device()
    assert reader.get('/api/knowledge/documents').json()['total'] == 0
    assert reader.get('/api/knowledge/search', params={'q': 'тест'}).json()['total'] == 0
    for path in [f"documents/{doc['id']}", f"documents/{doc['id']}/versions", f"versions/{version['id']}/sections", f"sections/{point['id']}", f"versions/{version['id']}/original"]:
        assert reader.get('/api/knowledge/' + path).status_code == 404
    assert invoke(reader, 'get_knowledge_section', {'sectionId': point['id']}).status_code == 404
    assert invoke(reader, 'verify_citation', {'documentId': doc['id'], 'versionId': version['id'], 'sectionId': point['id'], 'excerpt': 'test'}).status_code == 404
    assert invoke(reader, 'list_knowledge_versions', {'documentId': doc['id']}).status_code == 404
    assert invoke(reader, 'get_knowledge_document', {'documentId': doc['id']}).status_code == 404
    assert invoke(reader, 'search_knowledge', {'q': 'тест'}).json()['total'] == 0
    assert invoke(reader, 'list_knowledge_documents', {}).json()['total'] == 0
    assert upload(reader).status_code == 403
    assert reader.post(f"/api/knowledge/versions/{version['id']}/index").status_code == 403
    assert reader.get('/api/knowledge/audit').status_code == 403
    assert reader.post('/api/knowledge/admin-session', json={'token': 'wrong'}).status_code == 403
    anonymous = TestClient(app)
    assert anonymous.post('/api/knowledge/admin-session', json={'token': TOKEN}).status_code == 403
    assert invoke(anonymous, 'search_knowledge', {'q': 'test'}).status_code == 403
    assert admin.delete('/api/knowledge/admin-session').status_code == 200
    assert admin.get(version['originalUrl']).status_code == 404
    assert admin.get('/api/knowledge/session').json()['role'] == 'reader'


def test_upload_validation_safe_filenames_status_and_template_classification(corpus):
    admin = device(True)
    for content, filename, mime in [(b'%PDF bad', 'a.docx', DOCX_MIME), (b'fake', 'a.pdf', 'application/pdf'), (b'\x00\xff', 'a.txt', 'text/plain'), (b'test', 'a.exe', 'application/octet-stream'), (b'test', 'a.txt', 'image/png')]:
        assert upload(admin, content, filename, mime=mime).status_code == 422
    assert upload(admin, meta=metadata(status='active', effectiveFrom='2020-01-01')).status_code == 422
    assert upload(admin, meta=metadata(status='active', statusEvidence='test')).status_code == 422
    assert upload(admin, meta=metadata(title='   ')).status_code == 422
    assert upload(admin, b'too big' * (MAX_BYTES // 7 + 1)).status_code == 422
    created = upload(admin, name='../../escape.txt', meta=metadata(category='template')).json()
    assert created['latestVersion']['filename'] == 'escape.txt'
    assert not (corpus.root.parent / 'escape.txt').exists()
    assert admin.get('/api/knowledge/search', params={'q': 'тест'}).json()['total'] == 0
    assert admin.get('/api/knowledge/search', params={'q': 'тест', 'category': 'template'}).json()['total'] > 0
    doc = upload(admin, '1. Тестовая подтверждённая редакция'.encode(), meta=metadata(status='active', effectiveFrom='2020-01-01', statusEvidence='Основание только для автоматического теста')).json()
    assert admin.get('/api/knowledge/documents', params={'status': 'active'}).json()['total'] == 1
    assert admin.get('/api/knowledge/documents', params={'fileType': 'pdf'}).json()['total'] == 0
    assert doc['latestVersion']['statusEvidence']
    assert admin.get('/api/knowledge/versions/..%2F..%2Fdocuments.sqlite3/original').status_code in (404, 422)


def test_fragments_and_long_text_are_not_invented_numbering():
    extracted = extract('Текст без нормативного номера\nПродолжение'.encode(), 'txt', 'text/plain')
    assert extracted.sections[0].number is None
    assert extracted.sections[0].confidence == 'fragment'
    assert extracted.warnings
