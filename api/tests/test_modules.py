"""Real source-module regression checks for every added schema."""
import io
import re
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from api import main
from api.service import DocumentsService, SOURCE_DIR

IDS = ('project-committee-order', 'internal-audit-order', 'system-commissioning-order',
       'records-officers-order', 'handover-order', 'local-act-order')


def client_with_service(tmp_path, monkeypatch):
    monkeypatch.setattr(main, 'service', DocumentsService(tmp_path))
    client = TestClient(main.app)
    session = client.get('/api/session').json()
    return client, {'X-CSRF-Token': session['csrfToken']}


def sample_fields(specs):
    result = {}
    for spec in specs:
        if not spec.required: continue
        result[spec.key] = (spec.options[0] if spec.options else
                            '2026-12-31' if spec.inputType == 'date' else
                            f'Учебное значение {spec.key}')
    if 'appointees' in result:
        result['appointees'] = 'Управление документации | специалист И. И. Иванов\nУправление ИТ | аналитик П. П. Петров'
    return result


def text_of(content):
    document = Document(io.BytesIO(content))
    text = '\n'.join(p.text for p in document.paragraphs)
    text += '\n' + '\n'.join(c.text for t in document.tables for row in t.rows for c in row.cells)
    return document, text


def test_all_six_modules_generate_from_exact_supplied_source(tmp_path, monkeypatch):
    client, headers = client_with_service(tmp_path, monkeypatch)
    for template_id in IDS:
        template = client.get(f'/api/document-templates/{template_id}').json()
        assert template['sourceStatus'] == 'ready'
        assert template['registryStatus'] == 'unverified'
        source = (SOURCE_DIR / template['sourceFilename']).read_bytes()
        project = client.post('/api/document-projects', json={'templateId': template_id}, headers=headers).json()
        assert project['templateId'] == template_id
        fields = sample_fields(main.service.template(template_id).fields)
        preview = client.post(f"/api/document-projects/{project['projectId']}/preview", json={'fields': fields}, headers=headers)
        assert preview.status_code == 200, (template_id, preview.text)
        assert preview.json()['validation']['ready']
        assert preview.json()['preview']
        project = client.patch(f"/api/document-projects/{project['projectId']}", json={'expectedVersion': 1, 'fields': fields}, headers=headers).json()
        assert project['validation']['ready'], (template_id, project['validation'])
        generated = client.post(f"/api/document-projects/{project['projectId']}/generate", headers=headers)
        assert generated.status_code == 200, (template_id, generated.text)
        artifact = generated.json()['artifacts'][0]
        response = client.get(f"/api/document-projects/{project['projectId']}/artifacts/{artifact['artifactId']}/download")
        assert response.status_code == 200
        document, text = text_of(response.content)
        assert source != response.content
        assert 'ДАТА / РЕГИСТРАЦИОННЫЙ НОМЕР' in text
        assert 'ВИЗУАЛИЗАЦИЯ ЭП' in text
        assert 'МОДУЛЬНЫЙ ДОКУМЕНТ' not in text and 'УТВЕРЖДАЮ' not in text
        assert 'Примечания по заполнению' not in text
        assert not re.search(r'\[|\]|_{3,}|НЕ УДАЛЯТЬ', text)
        assert 'Учебное значение organization' in text
        assert len(document.tables) >= 2


def test_module_branches_and_field_validation(tmp_path, monkeypatch):
    client, headers = client_with_service(tmp_path, monkeypatch)
    committee = client.post('/api/document-projects', json={'templateId': 'project-committee-order'}, headers=headers).json()
    cid = committee['projectId']
    fields = sample_fields(main.service.template('project-committee-order').fields)
    fields.update(projectForm='проектам', projects='Проект А, Проект Б', cancelDate='2026-10-01',
                  cancelNumber='42', cancelTitle='О старом комитете')
    bad = {**fields, 'cancelNumber': ''}
    invalid = client.post(f'/api/document-projects/{cid}/preview', json={'fields': bad}, headers=headers).json()
    assert any(issue['code'] == 'conditional_required' for issue in invalid['validation']['issues'])
    project = client.patch(f'/api/document-projects/{cid}', json={'expectedVersion': 1, 'fields': fields}, headers=headers).json()
    artifact = client.post(f'/api/document-projects/{cid}/generate', headers=headers).json()['artifacts'][0]
    text = text_of(client.get(f"/api/document-projects/{cid}/artifacts/{artifact['artifactId']}/download").content)[1]
    assert 'проектам Проект А, Проект Б' in text
    assert '3. Контроль' in text and '2. Отменить распоряжение' in text
    assert '№ 42' in text and 'О старом комитете' in text
    assert project['validation']['ready']

    records = client.post('/api/document-projects', json={'templateId': 'records-officers-order'}, headers=headers).json()
    rid = records['projectId']
    fields = sample_fields(main.service.template('records-officers-order').fields)
    fields['appointees'] = 'без разделителя'
    invalid = client.post(f'/api/document-projects/{rid}/preview', json={'fields': fields}, headers=headers).json()
    assert any(issue['code'] == 'invalid_appointees' for issue in invalid['validation']['issues'])
    fields['appointees'] = '\n'.join(f'Управление {i} | специалист {i}' for i in range(4))
    fields['cancelOrder'] = 'приказ Организации от 1 октября 2026 г. № 42 «О назначении»'
    records = client.patch(f'/api/document-projects/{rid}', json={'expectedVersion': 1, 'fields': fields}, headers=headers).json()
    artifact = client.post(f'/api/document-projects/{rid}/generate', headers=headers).json()['artifacts'][0]
    document, text = text_of(client.get(f"/api/document-projects/{rid}/artifacts/{artifact['artifactId']}/download").content)
    assert len(document.tables[1].rows) == 5
    assert '5. Признать утратившим силу' in text and '6. Контроль' in text
    assert records['validation']['ready']

    audit = client.post('/api/document-projects', json={'templateId': 'internal-audit-order'}, headers=headers).json()
    aid = audit['projectId']
    fields = sample_fields(main.service.template('internal-audit-order').fields)
    fields.update(startDate='2026-12-31', endDate='2026-10-01')
    invalid = client.post(f'/api/document-projects/{aid}/preview', json={'fields': fields}, headers=headers).json()
    assert any(issue['code'] == 'date_order' for issue in invalid['validation']['issues'])

    handover = client.post('/api/document-projects', json={'templateId': 'handover-order'}, headers=headers).json()
    hid = handover['projectId']
    fields = sample_fields(main.service.template('handover-order').fields)
    fields['source'] = 'от начальника Управления'
    client.patch(f'/api/document-projects/{hid}', json={'expectedVersion': 1, 'fields': fields}, headers=headers)
    artifact = client.post(f'/api/document-projects/{hid}/generate', headers=headers).json()['artifacts'][0]
    text = text_of(client.get(f"/api/document-projects/{hid}/artifacts/{artifact['artifactId']}/download").content)[1]
    assert 'от начальника Управления' in text and 'от от начальника' not in text
