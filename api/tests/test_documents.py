import io
import os
import tempfile
from pathlib import Path

os.environ['REGULA_DATA_DIR'] = tempfile.mkdtemp(prefix='regula-tests-')
os.environ['REGULA_SESSION_SECRET'] = 'test-secret-long-enough'

from fastapi.testclient import TestClient
from docx import Document
from api.main import app
from api.service import FIELDS, SOURCE_DIR


def device():
    client = TestClient(app)
    response = client.get('/api/session')
    assert response.status_code == 200
    return client, {'X-CSRF-Token': response.json()['csrfToken']}


def values():
    return {spec.key: (f'Тестовое значение {spec.key}' if spec.key != 'deadline' else '2026-12-31') for spec in FIELDS}


def test_catalog_project_lifecycle_and_device_isolation():
    client, headers = device()
    catalog = client.get('/api/document-templates').json()
    assert len(catalog) == 7
    assert catalog[0]['sourceStatus'] == 'ready'
    assert catalog[0]['registryStatus'] == 'unverified'
    assert all(item['sourceStatus'] == 'ready' for item in catalog)
    other_template = client.post('/api/document-projects', json={'templateId': catalog[1]['templateId']}, headers=headers)
    assert other_template.status_code == 200 and not other_template.json()['validation']['ready']
    created = client.post('/api/document-projects', json={'templateId': catalog[0]['templateId']}, headers=headers)
    assert created.status_code == 200
    project = created.json()
    project_id = project['projectId']
    assert project['validation']['ready'] is False
    assert client.post(f'/api/document-projects/{project_id}/generate', headers=headers).status_code == 422
    assert client.post(f'/api/document-projects/{project_id}/preview', json={'fields': values()}, headers=headers).json()['validation']['ready']
    changed = client.patch(f'/api/document-projects/{project_id}', json={'expectedVersion': 1, 'fields': values()}, headers=headers)
    assert changed.status_code == 200
    assert changed.json()['version'] == 2
    assert changed.json()['validation']['ready'] is True
    assert client.patch(f'/api/document-projects/{project_id}', json={'expectedVersion': 1, 'fields': {'chair': 'Иное'}}, headers=headers).status_code == 409
    assert client.get('/api/document-projects').json()[0]['projectId'] == project_id
    tools = client.get('/api/document-tools').json()
    assert tools[2]['inputSchema']['properties']['templateId']
    assert client.post('/api/document-tools/create_document_project/invoke', json={'arguments': {'templateId': 'working-group-order'}}).status_code == 403
    invoked = client.post('/api/document-tools/get_document_project/invoke', json={'arguments': {'projectId': project_id}}, headers=headers)
    assert invoked.status_code == 200 and invoked.json()['projectId'] == project_id
    other, other_headers = device()
    assert other.get('/api/document-projects').json() == []
    assert other.get(f'/api/document-projects/{project_id}').status_code == 404
    assert other.patch(f'/api/document-projects/{project_id}', json={'expectedVersion': 2, 'fields': {}}, headers=other_headers).status_code == 404
    assert other.post(f'/api/document-projects/{project_id}/generate', headers=other_headers).status_code == 404
    assert other.get(f'/api/document-projects/{project_id}/versions').status_code == 404
    assert client.get('/api/document-projects/../data/documents.sqlite3').status_code in (404, 405)


def test_real_source_docx_preserves_service_areas_and_versions():
    client, headers = device()
    project = client.post('/api/document-projects', json={'templateId': 'working-group-order'}, headers=headers).json()
    project_id = project['projectId']
    project = client.patch(f'/api/document-projects/{project_id}', json={'expectedVersion': 1, 'fields': values()}, headers=headers).json()
    generated = client.post(f'/api/document-projects/{project_id}/generate', headers=headers)
    assert generated.status_code == 200, generated.text
    assert len(client.post(f'/api/document-projects/{project_id}/generate', headers=headers).json()['artifacts']) == 1
    unchanged = client.patch(f'/api/document-projects/{project_id}', json={'expectedVersion': 2, 'fields': values()}, headers=headers).json()
    assert unchanged['version'] == 2 and unchanged['status'] == 'generated'
    artifact = generated.json()['artifacts'][0]
    download = client.get(f'/api/document-projects/{project_id}/artifacts/{artifact["artifactId"]}/download')
    assert download.status_code == 200
    document = Document(io.BytesIO(download.content))
    text = '\n'.join(p.text for p in document.paragraphs)
    text += '\n' + '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    for key in ('purpose', 'subject', 'chair', 'secretary', 'members', 'tasks', 'result', 'recipient', 'termination', 'control', 'signer'):
        assert f'Тестовое значение {key}' in text
    assert '31» декабря 2026 г.' in text
    assert 'ДАТА / РЕГИСТРАЦИОННЫЙ НОМЕР' in text
    assert 'ВИЗУАЛИЗАЦИЯ ЭП' in text
    assert 'МОДУЛЬНЫЙ ДОКУМЕНТ' not in text
    assert 'УТВЕРЖДАЮ' not in text
    assert '[1]' not in text and '___' not in text
    assert len(document.tables) == 3
    assert client.get(f'/api/document-projects/{project_id}/versions').json()[0]['artifactId'] == artifact['artifactId']
    other, _ = device()
    assert other.get(f'/api/document-projects/{project_id}/artifacts/{artifact["artifactId"]}/download').status_code == 404
    assert client.get(f'/api/document-projects/{project_id}/artifacts/..%2F..%2Fdocuments.sqlite3/download').status_code in (404, 422)
    changed = client.patch(f'/api/document-projects/{project_id}', json={'expectedVersion': 2, 'fields': {'subject': 'Новая тема'}}, headers=headers).json()
    assert changed['status'] == 'draft'
    second = client.post(f'/api/document-projects/{project_id}/generate', headers=headers).json()
    assert len(second['artifacts']) == 2
    assert second['artifacts'][0]['projectVersion'] == 3
    assert (SOURCE_DIR/'working-group-order.docx').read_bytes() != download.content


def test_anonymous_upload_is_disabled_and_invalid_data_rejected():
    client, headers = device()
    source = (SOURCE_DIR/'working-group-order.docx').read_bytes()
    upload = client.post('/api/document-templates/working-group-order/source', headers=headers, files={'file': ('module.docx', source, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert upload.status_code == 403
    assert client.post('/api/document-templates/working-group-order/approve', headers=headers).status_code == 403
    project = client.post('/api/document-projects', json={'templateId': 'working-group-order'}, headers=headers).json()
    project_id = project['projectId']
    assert client.patch(f'/api/document-projects/{project_id}', json={'expectedVersion': 1, 'fields': {'bad': 'x'}}, headers=headers).status_code == 422
    assert client.patch(f'/api/document-projects/{project_id}', json={'expectedVersion': 1, 'fields': {'purpose': 'x'*4001}}, headers=headers).status_code == 422
    assert client.post(f'/api/document-projects/{project_id}/preview', json={'fields': {'bad': 'x'}}, headers=headers).status_code == 422
    assert client.get(f'/api/document-projects/{project_id}/artifacts/not-a-uuid/download').status_code == 404
