"""Checked, source-based renderers for the six remaining synthetic DOCX modules.

Each renderer changes a fresh in-memory copy and fails when expected anchors move.
Official status, signing, registration and IIS DO actions are outside this service.
"""
import re
from datetime import date
from io import BytesIO

from docx import Document

from .working_group import MONTHS, drop_paragraph, drop_table, replace, replace_pattern

MODULES = {
    'project-committee-order': ('О создании проектного комитета', 3, 13),
    'internal-audit-order': ('О проведении внутреннего аудита', 4, 19),
    'system-commissioning-order': ('О вводе в эксплуатацию', 3, 16),
    'records-officers-order': ('О назначении ответственных', 4, 18),
    'handover-order': ('О приеме и передаче', 3, 12),
    'local-act-order': ('Об утверждении и введении', 3, 16),
}


def value(fields, key, draft):
    return fields.get(key, '').strip() or ('…' if draft else '')


def date_value(fields, key, draft):
    raw = fields.get(key, '').strip()
    if not raw:
        return '…' if draft else ''
    try:
        parsed = date.fromisoformat(raw)
    except ValueError:
        if draft:
            return '…'
        raise
    return f'«{parsed.day:02d}» {MONTHS[parsed.month-1]} {parsed.year} г.'


def slot(paragraph, new, pattern=r'_{3,}'):
    replace_pattern(paragraph, pattern, new)


def exact(paragraph, old, new):
    if old not in paragraph.text:
        raise ValueError('source_structure_changed')
    replace(paragraph, old, new)


def fill_signature(document, fields, draft):
    table = document.tables[-1]
    if 'И. О. Фамилия' not in table.cell(0, 2).text:
        raise ValueError('source_structure_changed')
    replace_pattern(table.cell(0, 2).paragraphs[0], r'И\. О\. Фамилия\s*\[\d+\]', value(fields, 'signer', draft))
    slot(table.cell(0, 1).paragraphs[0], '')
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True


def fill_committee(p, document, fields, draft):
    form = value(fields, 'projectForm', draft)
    plural = form == 'проектам'
    for index in (5, 8):
        exact(p[index], '[проекту / проектам]', form)
        slot(p[index], value(fields, 'projects', draft))
    slot(p[6], value(fields, 'purpose', draft))
    if fields.get('reason', '').strip():
        slot(p[6], value(fields, 'reason', draft))
    else:
        replace_pattern(p[6], r'в связи с _{3,}\s*\[4\],\s*', '')
    exact(p[6], '[проекта / проектов]', 'проектов' if plural else 'проекта')
    slot(p[6], value(fields, 'projects', draft))
    exact(p[6], 'реализуемого(ых)', 'реализуемых' if plural else 'реализуемого')
    slot(p[6], value(fields, 'implementingDepartment', draft))
    if fields.get('cancelDate') or fields.get('cancelNumber') or fields.get('cancelTitle'):
        old_date = date_value(fields, 'cancelDate', draft)
        number = value(fields, 'cancelNumber', draft)
        title = value(fields, 'cancelTitle', draft)
        replace_pattern(p[5], r'\[и об отмене распоряжения Организации от [^\]]+\]',
                        f'и об отмене распоряжения Организации от {old_date} № {number}')
        replace_pattern(p[9], r'\[Отменить распоряжение Организации от [^\]]+\]',
                        f'Отменить распоряжение Организации от {old_date} № {number} «{title}».')
    else:
        replace_pattern(p[5], r'\s*\[и об отмене распоряжения Организации от [^\]]+\]', '')
        drop_paragraph(p[9])
        exact(p[10], '3. Контроль', '2. Контроль')
    slot(p[10], value(fields, 'control', draft))
    composition = document.tables[2]
    roles = ['Председатель Комитета', 'Заместитель председателя Комитета', 'Секретарь Комитета', 'Члены Комитета']
    for row, (role, key) in enumerate(zip(roles, ('chair', 'viceChair', 'secretary', 'members')), 1):
        if composition.cell(row, 0).text.strip() != role:
            raise ValueError('source_structure_changed')
        slot(composition.cell(row, 1).paragraphs[0], value(fields, key, draft))


def fill_audit(p, document, fields, draft):
    form = value(fields, 'processForm', draft)
    for index in (5, 8):
        exact(p[index], 'процесса / процессов', form)
    slot(p[5], value(fields, 'processes', draft))
    slot(p[6], value(fields, 'basis', draft))
    replace_pattern(p[8], r'«___»\s*_{3,}\s*по\s*«___»\s*_{3,}\s*20___\s*г\.',
                    f"{date_value(fields, 'startDate', draft)} по {date_value(fields, 'endDate', draft)}")
    slot(p[8], value(fields, 'processes', draft))
    slot(p[14], value(fields, 'auditedDepartment', draft))
    replace_pattern(p[15], r'«___»\s*_{3,}\s*20___\s*г\.', date_value(fields, 'reportDate', draft))
    slot(p[15], value(fields, 'reportRecipient', draft))
    slot(p[16], value(fields, 'actionOwner', draft))
    replace_pattern(p[16], r'«___»\s*_{3,}\s*20___\s*г\.', date_value(fields, 'actionDate', draft))
    slot(p[17], value(fields, 'control', draft))
    composition = document.tables[2]
    for row, role, key in ((1, 'Руководитель', 'chair'), (2, 'Аудиторы', 'auditors')):
        if composition.cell(row, 0).text.strip() != role:
            raise ValueError('source_structure_changed')
        slot(composition.cell(row, 1).paragraphs[0], value(fields, key, draft))
    fill_signature(document, fields, draft)


def fill_system(p, document, fields, draft):
    slot(p[5], value(fields, 'systemName', draft))
    slot(p[6], value(fields, 'readinessDocument', draft))
    slot(p[6], value(fields, 'systemName', draft))
    replace_pattern(p[6], r'\(далее - _{3,}»\)', f"(далее - «{value(fields, 'shortName', draft)}»)")
    exact(p[6], 'разработанного / измененного', value(fields, 'developmentMode', draft))
    slot(p[6], value(fields, 'workDocument', draft))
    slot(p[8], value(fields, 'shortName', draft))
    replace_pattern(p[8], r'«___»\s*_{3,}\s*20___\s*г\.', date_value(fields, 'effectiveDate', draft))
    slot(p[9], value(fields, 'businessOwner', draft))
    slot(p[10], value(fields, 'shortName', draft))
    slot(p[10], value(fields, 'purpose', draft))
    exact(p[11], 'Департаменту информационных технологий', value(fields, 'itDepartment', draft))
    for index in (12, 13):
        slot(p[index], value(fields, 'shortName', draft))
    slot(p[14], value(fields, 'responsible', draft))
    fill_signature(document, fields, draft)


def fill_appointees(table, fields, draft):
    lines = [line.strip() for line in fields.get('appointees', '').splitlines() if line.strip()]
    if not lines and draft:
        lines = ['… | …']
    if not lines or any('|' not in line for line in lines):
        raise ValueError('invalid_appointees')
    pairs = [tuple(part.strip() for part in line.split('|', 1)) for line in lines]
    if any(not left or not right for left, right in pairs):
        raise ValueError('invalid_appointees')
    if table.cell(0, 0).text.strip() != 'Подразделение':
        raise ValueError('source_structure_changed')
    while len(table.rows) < len(pairs) + 1:
        table.add_row()
    while len(table.rows) > len(pairs) + 1:
        table._tbl.remove(table.rows[-1]._tr)
    for row, (department, person) in enumerate(pairs, 1):
        cell0, cell1 = table.rows[row].cells
        if re.search(r'_{3,}', cell0.text):
            slot(cell0.paragraphs[0], department)
        else:
            cell0.text = department
        if re.search(r'_{3,}', cell1.text):
            slot(cell1.paragraphs[0], person)
        else:
            cell1.text = person


def fill_records(p, document, fields, draft):
    slot(p[6], value(fields, 'basis', draft))
    slot(p[13], value(fields, 'registryDepartment', draft))
    slot(p[14], value(fields, 'coordinator', draft))
    if fields.get('cancelOrder', '').strip():
        replace_pattern(p[15], r'\[Признать утратившим силу _{3,}\.\]',
                        f"Признать утратившим силу {value(fields, 'cancelOrder', draft)}.")
    else:
        drop_paragraph(p[15])
        exact(p[16], '6. Контроль', '5. Контроль')
    slot(p[16], value(fields, 'control', draft))
    p[16].paragraph_format.keep_with_next = True
    fill_appointees(document.tables[2], fields, draft)
    fill_signature(document, fields, draft)


def fill_handover(p, document, fields, draft):
    slot(p[6], value(fields, 'reason', draft))
    for key in ('giver', 'receiver', 'items'):
        slot(p[8], value(fields, key, draft))
    if fields.get('source', '').strip():
        slot(p[8], re.sub(r'^от\s+', '', value(fields, 'source', draft), flags=re.IGNORECASE))
    else:
        replace_pattern(p[8], r'\s*от _{3,}\s*\[5\]', '')
    exact(p[8], ' в порядке,', f" не позднее {date_value(fields, 'deadline', draft)} в порядке,")
    slot(p[9], value(fields, 'representative', draft))
    slot(p[10], value(fields, 'control', draft))
    fill_signature(document, fields, draft)


def fill_local_act(p, document, fields, draft):
    slot(p[5], value(fields, 'actName', draft))
    slot(p[6], value(fields, 'purpose', draft))
    slot(p[8], value(fields, 'actName', draft))
    replace_pattern(p[9], r'«___»\s*_{3,}\s*20___\s*г\.', date_value(fields, 'effectiveDate', draft))
    slot(p[10], value(fields, 'responsibleDepartment', draft))
    slot(p[12], value(fields, 'publicationPlace', draft))
    if fields.get('cancelOrder', '').strip():
        replace_pattern(p[13], r'\[Признать утратившим силу _{3,}\.\]',
                        f"Признать утратившим силу {value(fields, 'cancelOrder', draft)}.")
    else:
        drop_paragraph(p[13])
        exact(p[14], '5. Контроль', '4. Контроль')
    slot(p[14], value(fields, 'control', draft))
    fill_signature(document, fields, draft)

FILLERS = {
    'project-committee-order': fill_committee,
    'internal-audit-order': fill_audit,
    'system-commissioning-order': fill_system,
    'records-officers-order': fill_records,
    'handover-order': fill_handover,
    'local-act-order': fill_local_act,
}


def render(template_id: str, source: bytes, fields: dict[str, str], *, draft=False) -> bytes:
    title, table_count, notes_start = MODULES[template_id]
    document = Document(BytesIO(source))
    p = document.paragraphs
    if len(p) <= notes_start or len(document.tables) != table_count:
        raise ValueError('source_structure_changed')
    if p[0].text != 'МОДУЛЬНЫЙ ДОКУМЕНТ' or p[1].text != 'НАИМЕНОВАНИЕ ОРГАНИЗАЦИИ':
        raise ValueError('source_structure_changed')
    if not p[5].text.startswith(title) or not p[notes_start].text.startswith('Примечания по заполнению'):
        raise ValueError('source_structure_changed')
    if 'УТВЕРЖДАЮ' not in document.tables[0].cell(0, 1).text or 'ГЕРБ' not in document.tables[1].cell(0, 0).text:
        raise ValueError('source_structure_changed')
    if 'ДАТА / РЕГИСТРАЦИОННЫЙ НОМЕР' not in p[3].text:
        raise ValueError('source_structure_changed')
    exact(p[1], 'НАИМЕНОВАНИЕ ОРГАНИЗАЦИИ', value(fields, 'organization', draft))
    FILLERS[template_id](p, document, fields, draft)
    # Remove module-only material, leaving the service areas intact.
    for paragraph in p[notes_start:]:
        drop_paragraph(paragraph)
    drop_paragraph(p[0])
    drop_table(document.tables[0])
    for paragraph in document.tables[0].cell(0, 0).paragraphs:
        replace(paragraph, 'ГЕРБ', '')
        replace(paragraph, 'место для изображения', '')
    for paragraph in document.paragraphs:
        replace(paragraph, ' (НЕ УДАЛЯТЬ)', '')
        for hint in re.findall(r'\s*\[\d+\]', paragraph.text):
            replace(paragraph, hint, '')
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace(paragraph, ' (НЕ УДАЛЯТЬ)', '')
                    for hint in re.findall(r'\s*\[\d+\]', paragraph.text):
                        replace(paragraph, hint, '')
    text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
    text += '\n' + '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    if re.search(r'\[|\]|_{3,}|МОДУЛЬНЫЙ ДОКУМЕНТ|УТВЕРЖДАЮ|НЕ УДАЛЯТЬ|Примечания по заполнению', text):
        raise ValueError('unresolved_placeholders')
    if 'ДАТА / РЕГИСТРАЦИОННЫЙ НОМЕР' not in text or 'ВИЗУАЛИЗАЦИЯ ЭП' not in text:
        raise ValueError('service_area_missing')
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def preview(template_id: str, source: bytes, fields: dict[str, str]):
    content = render(template_id, source, fields, draft=True)
    document = Document(BytesIO(content))
    items = []
    first_labels = {0: 'Организация', 1: 'Вид документа', 3: 'Место', 4: 'Заголовок',
                    5: 'Основание', 6: 'Распорядительная часть'}
    for index, paragraph in enumerate(document.paragraphs):
        content = paragraph.text.strip()
        if content and content not in ('ДАТА / РЕГИСТРАЦИОННЫЙ НОМЕР', 'ВИЗУАЛИЗАЦИЯ ЭП'):
            point = re.match(r'^(\d+(?:\.\d+)?)\.', content)
            label = first_labels.get(index) or (f'Пункт {point.group(1)}' if point else 'Текст документа')
            items.append({'section': 'Текст', 'label': label, 'value': content})
    for table in document.tables:
        if len(table.rows) > 2:
            for row in table.rows[1:]:
                items.append({'section': 'Состав', 'label': row.cells[0].text.strip(), 'value': row.cells[1].text.strip()})
    return items
